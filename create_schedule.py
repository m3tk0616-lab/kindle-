import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "スケジュール2026"

# Styles
header_font = Font(bold=True, size=11)
cell_font = Font(size=10)
center = Alignment(horizontal='center', vertical='center')
thin = Side(style='thin')
medium = Side(style='medium')
thin_border = Border(left=thin, right=thin, top=thin, bottom=thin)
medium_border = Border(left=medium, right=medium, top=medium, bottom=medium)

header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
month_fill = PatternFill(start_color="D6E4F7", end_color="D6E4F7", fill_type="solid")
sub_header_fill = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
light_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
holiday_fill = PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid")

# Data
# Format: (day, weekday, match, referee)
may_data = [
    (1, "金", "", ""),
    (2, "土", "", ""),
    (3, "日", "", ""),
    (4, "月", "", ""),
    (5, "火", "", ""),
    (6, "水", "", ""),
    (7, "木", "", ""),
    (8, "金", "", ""),
    (9, "土", "", ""),
    (10, "日", "", ""),
    (11, "月", "", ""),
    (12, "火", "", ""),
    (13, "水", "", ""),
    (14, "木", "", ""),
    (15, "金", "", ""),
    (16, "土", "", ""),
    (17, "日", "", ""),
    (18, "月", "", ""),
    (19, "火", "", ""),
    (20, "水", "", ""),
    (21, "木", "", ""),
    (22, "金", "", ""),
    (23, "土", "", ""),
    (24, "日", "", ""),
    (25, "月", "", ""),
    (26, "火", "1-3", "2-6"),
    (27, "水", "5-6", "1-3"),
    (28, "木", "2-4", "5-6"),
    (29, "金", "", ""),
    (30, "土", "", ""),
    (31, "日", "", ""),
]

june_data = [
    (1, "月", "", ""),
    (2, "火", "1-6", "2-4"),
    (3, "水", "2-3", "1-6"),
    (4, "木", "4-5", "2-3"),
    (5, "金", "", ""),
    (6, "土", "", ""),
    (7, "日", "", ""),
    (8, "月", "", ""),
    (9, "火", "1-2", "4-5"),
    (10, "水", "3-5", "1-2"),
    (11, "木", "4-6", "3-5"),
    (12, "金", "", ""),
    (13, "土", "", ""),
    (14, "日", "", ""),
    (15, "月", "", ""),
    (16, "火", "予備日", ""),
    (17, "水", "予備日", ""),
    (18, "木", "予備日", ""),
    (19, "金", "", ""),
    (20, "土", "", ""),
    (21, "日", "", ""),
    (22, "月", "", ""),
    (23, "火", "2-5", "4-6"),
    (24, "水", "1-4", "2-5"),
    (25, "木", "3-6", "1-4"),
    (26, "金", "", ""),
    (27, "土", "", ""),
    (28, "日", "", ""),
    (29, "月", "", ""),
    (30, "火", "1-5", "3-6"),
]

july_data = [
    (1, "水", "3-4", "1-5"),
    (2, "木", "2-6", "3-4"),
    (3, "金", "", ""),
    (4, "土", "", ""),
    (5, "日", "", ""),
    (6, "月", "", ""),
    (7, "火", "", ""),
    (8, "水", "", ""),
    (9, "木", "", ""),
    (10, "金", "", ""),
    (11, "土", "", ""),
    (12, "日", "", ""),
    (13, "月", "", ""),
    (14, "火", "", ""),
    (15, "水", "", ""),
    (16, "木", "", ""),
    (17, "金", "", ""),
    (18, "土", "", ""),
    (19, "日", "", ""),
    (20, "月", "", ""),
    (21, "火", "", ""),
    (22, "水", "", ""),
    (23, "木", "", ""),
    (24, "金", "", ""),
    (25, "土", "", ""),
    (26, "日", "", ""),
    (27, "月", "", ""),
    (28, "火", "", ""),
    (29, "水", "", ""),
    (30, "木", "", ""),
    (31, "金", "", ""),
]

# Column layout: 
# May:   A=日, B=曜日, C=対戦チーム, D=審判
# June:  F=日, G=曜日, H=対戦チーム, I=審判
# July:  K=日, L=曜日, M=対戦チーム, N=審判
# Spacer columns: E, J

col_offsets = [1, 6, 11]  # A=1, F=6, K=11
months = ["5月", "6月", "7月"]
all_data = [may_data, june_data, july_data]
sub_headers = ["日", "曜日", "対戦チーム", "審判"]

# Set column widths
col_widths = {
    1: 5,   # A 日
    2: 6,   # B 曜日
    3: 12,  # C 対戦チーム
    4: 8,   # D 審判
    5: 2,   # E spacer
    6: 5,   # F 日
    7: 6,   # G 曜日
    8: 12,  # H 対戦チーム
    9: 8,   # I 審判
    10: 2,  # J spacer
    11: 5,  # K 日
    12: 6,  # L 曜日
    13: 12, # M 対戦チーム
    14: 8,  # N 審判
}
for col_num, width in col_widths.items():
    ws.column_dimensions[get_column_letter(col_num)].width = width

# Row 1: Month headers (merged)
for i, (col_off, month) in enumerate(zip(col_offsets, months)):
    start_col = col_off
    end_col = col_off + 3
    ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)
    cell = ws.cell(row=1, column=start_col, value=month)
    cell.font = Font(bold=True, size=13, color="FFFFFF")
    cell.alignment = center
    cell.fill = header_fill
    cell.border = thin_border
    # Also apply border to merged cells
    for c in range(start_col, end_col+1):
        ws.cell(row=1, column=c).border = thin_border

# Row 2: Sub-headers
for col_off in col_offsets:
    for j, header in enumerate(sub_headers):
        cell = ws.cell(row=2, column=col_off + j, value=header)
        cell.font = Font(bold=True, size=10)
        cell.alignment = center
        cell.fill = sub_header_fill
        cell.border = thin_border

# Data rows (row 3 onwards)
weekday_colors = {
    "土": PatternFill(start_color="CCE5FF", end_color="CCE5FF", fill_type="solid"),  # blue
    "日": PatternFill(start_color="FFE0E0", end_color="FFE0E0", fill_type="solid"),  # red
}

for month_idx, (col_off, data) in enumerate(zip(col_offsets, all_data)):
    for day_idx, (day, weekday, match, referee) in enumerate(data):
        row = day_idx + 3
        values = [day, weekday, match, referee]
        
        # Determine fill
        if weekday == "土":
            fill = weekday_colors["土"]
        elif weekday == "日":
            fill = weekday_colors["日"]
        elif day_idx % 2 == 0:
            fill = white_fill
        else:
            fill = light_fill
        
        for j, val in enumerate(values):
            cell = ws.cell(row=row, column=col_off + j, value=val)
            cell.font = cell_font
            cell.alignment = center
            cell.fill = fill
            cell.border = thin_border
            
            # Highlight match/referee cells with data
            if j in (2, 3) and val and val != "予備日":
                cell.font = Font(size=10, bold=True, color="1F4E79")

# Set row heights
ws.row_dimensions[1].height = 22
ws.row_dimensions[2].height = 18
for r in range(3, 35):
    ws.row_dimensions[r].height = 16

# Freeze pane
ws.freeze_panes = "A3"

output_path = "/home/user/kindle-/schedule_2026.xlsx"
wb.save(output_path)
print(f"Saved: {output_path}")
