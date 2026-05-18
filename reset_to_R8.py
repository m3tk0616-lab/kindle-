import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from copy import copy
import datetime

SRC = "/root/.claude/uploads/5a98ff47-aadd-49b8-9dad-a8d2f564f977/3ab50f35-__7________.xlsx"
DST = "/home/user/kindle-/schedule_R8.xlsx"

wb = openpyxl.load_workbook(SRC)

def safe_set(ws, row, col, value):
    """Set cell value, skipping merged (read-only) cells."""
    cell = ws.cell(row, col)
    if isinstance(cell, MergedCell):
        return False
    cell.value = value
    return True

# ════════════════════════════════════════════
# リーグ勝敗表 — リセット & 令和8年チーム差替
# ════════════════════════════════════════════
ws = wb["リーグ勝敗表"]

# タイトル追加 (A1:S1 を結合して題名)
ws.merge_cells("A1:S1")
title_cell = ws.cell(1, 1)
title_cell.value = "令和8年 朝野球リーグ勝敗表"
title_cell.font = Font(name="メイリオ", size=16, bold=True)
title_cell.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 25

# 1) 入力スコアを全クリア (E列〜AK列, 数値のみ — 数式は触らない)
for row in ws.iter_rows(min_row=6, max_row=27, min_col=5, max_col=37):
    for cell in row:
        if isinstance(cell, MergedCell):
            continue
        if isinstance(cell.value, (int, float)):
            cell.value = None

# 2) チームリスト (BA列) を令和8年6チームに差替
new_teams = ["野宮塗装", "飯詰", "バッカス", "ウエスタン", "富士電機", "日産"]
for i, name in enumerate(new_teams):
    safe_set(ws, 6+i, 53, name)  # BA = col 53
# 7番以降はクリア
for r in range(12, 21):
    safe_set(ws, r, 53, None)

# 3) 出場チーム番号 (AW列=49) を 1〜6 に
team_slots = [6, 8, 10, 12, 14, 16, 18, 20, 22, 24, 26]
for i, r in enumerate(team_slots):
    safe_set(ws, r, 49, (i + 1) if i < 6 else None)

# ════════════════════════════════════════════
# 日程表１ — 年度・曜日・対戦表を令和8年に
# ════════════════════════════════════════════
ws1 = wb["日程表１"]
safe_set(ws1, 1, 1, "R8朝野球日程表")
safe_set(ws1, 1, 20, 2026)

# 曜日列を2026年で再計算
WD = ["月","火","水","木","金","土","日"]
weekday_setup = [(2, 5), (8, 6), (14, 7), (20, 8)]  # (列, 月)

for row in range(4, 35):
    day = row - 3
    for col, month in weekday_setup:
        try:
            d = datetime.date(2026, month, day)
            safe_set(ws1, row, col, WD[d.weekday()])
        except ValueError:
            safe_set(ws1, row, col, None)

# 対戦・審判列
match_cols = {
    5: (3, 5, 6),    # 5月: C, E, F
    6: (9, 11, 12),  # 6月: I, K, L
    7: (15, 17, 18), # 7月: O, Q, R
    8: (21, 23, 24), # 8月: U, W, X
}
sep_cols = {5: 4, 6: 10, 7: 16, 8: 22}  # ー separator cols

# 既存の対戦データを全クリア
for month, (ca, cb, cr) in match_cols.items():
    for row in range(4, 35):
        safe_set(ws1, row, ca, None)
        safe_set(ws1, row, cb, None)
        safe_set(ws1, row, cr, None)

# 令和8年 PDF スケジュール
GAMES = {
    (5,26): (1, 3, "2-6"),
    (5,27): (5, 6, "1-3"),
    (5,28): (2, 4, "5-6"),
    (6, 2): (1, 6, "2-4"),
    (6, 3): (2, 3, "1-6"),
    (6, 4): (4, 5, "2-3"),
    (6, 9): (1, 2, "4-5"),
    (6,10): (3, 5, "1-2"),
    (6,11): (4, 6, "3-5"),
    (6,16): "予備日",
    (6,17): "予備日",
    (6,18): "予備日",
    (6,23): (2, 5, "4-6"),
    (6,24): (1, 4, "2-5"),
    (6,25): (3, 6, "1-4"),
    (6,30): (1, 5, "3-6"),
    (7, 1): (3, 4, "1-5"),
    (7, 2): (2, 6, "3-4"),
}

LKP = "リーグ勝敗表!$AZ$6:$BA$11"

def team_formula(n):
    return f'=IFERROR(VLOOKUP({n},{LKP},2,0),"{n}")'

def ref_formula(ref_str):
    """'2-6' → VLOOKUP(2)&'-'&VLOOKUP(6)"""
    parts = ref_str.split("-")
    if len(parts) == 2:
        a, b = parts
        return (f'=IFERROR(VLOOKUP({a},{LKP},2,0),"{a}")'
                f'&"-"&IFERROR(VLOOKUP({b},{LKP},2,0),"{b}")')
    return ref_str

def write_match(ws, row, month, data):
    ca, cb, cr = match_cols[month]
    sc = sep_cols[month]
    if data == "予備日":
        safe_set(ws, row, ca, "予備日")
        safe_set(ws, row, sc, None)
        safe_set(ws, row, cb, None)
        safe_set(ws, row, cr, None)
    else:
        ta, tb, ref = data
        safe_set(ws, row, ca, team_formula(ta))
        safe_set(ws, row, cb, team_formula(tb))
        safe_set(ws, row, cr, ref_formula(ref))

for (month, day), data in GAMES.items():
    write_match(ws1, day + 3, month, data)

# ════════════════════════════════════════════
# 日程表２ — 同じスケジュールを反映
# ════════════════════════════════════════════
ws2 = wb["日程表２"]
for month, (ca, cb, cr) in match_cols.items():
    for row in range(4, 35):
        safe_set(ws2, row, ca, None)
        safe_set(ws2, row, cb, None)
        safe_set(ws2, row, cr, None)

for (month, day), data in GAMES.items():
    write_match(ws2, day + 3, month, data)

# ════════════════════════════════════════════
# 文字切れ対策：横向き・縮小して全体表示・列幅調整
# ════════════════════════════════════════════
def shrink_match_cell(cell, sz=8):
    if isinstance(cell, MergedCell):
        return
    f = cell.font
    cell.font = Font(name=f.name or "メイリオ", size=sz,
                     bold=f.bold, italic=f.italic, color=f.color)
    a = cell.alignment
    cell.alignment = Alignment(
        horizontal=a.horizontal or "center",
        vertical=a.vertical or "center",
        shrink_to_fit=True,
        wrap_text=False,
    )

# 各シートの対戦/審判セル全部に縮小フィット適用
match_all_cols = []
for ca, cb, cr in match_cols.values():
    match_all_cols.extend([ca, cb, cr])

for ws_x in (ws1, ws2):
    # ページ設定：A4横、1ページに収める
    ws_x.page_setup.orientation = "landscape"
    ws_x.page_setup.paperSize = 9
    ws_x.page_setup.fitToWidth = 1
    ws_x.page_setup.fitToHeight = 1
    ws_x.sheet_properties.pageSetUpPr.fitToPage = True
    ws_x.page_margins.left = 0.3
    ws_x.page_margins.right = 0.3
    ws_x.page_margins.top = 0.4
    ws_x.page_margins.bottom = 0.4

    # 対戦・審判セルを小フォント＋縮小フィット
    for row in range(4, 35):
        for col in match_all_cols:
            shrink_match_cell(ws_x.cell(row, col), sz=8)

    # 細すぎる team A 列（C/I/O/U = 3/9/15/21）を広げる
    for col_letter in ["C", "I", "O", "U"]:
        ws_x.column_dimensions[col_letter].width = 8
    # 区切り "ー" 列（D/J/P/V）は狭く
    for col_letter in ["D", "J", "P", "V"]:
        ws_x.column_dimensions[col_letter].width = 2.5
    # 審判列（F/L/R/X）を少し広げる
    for col_letter in ["F", "L", "R", "X"]:
        ws_x.column_dimensions[col_letter].width = 11

# ════════════════════════════════════════════
# 既存の網掛け(黄色)を除去 & 飯詰試合日に網掛け
# ════════════════════════════════════════════
NO_FILL = PatternFill(fill_type=None)
# 元ファイルで使われていた黄色 FFFFFF00 を飯詰部分に流用
IIZUME_FILL = PatternFill(patternType="solid",
                          fgColor="FFFFFF00", bgColor="FFFFFF00")

# 飯詰(2)が試合に出る日 → 月別 (day, 月)
iizume_days_by_month = {5: [28], 6: [3, 9, 23], 7: [2]}

def clear_all_fills(ws):
    """データ域(rows 4-34, cols A-X)の全fillを除去"""
    for row in ws.iter_rows(min_row=4, max_row=34, min_col=1, max_col=24):
        for c in row:
            if isinstance(c, MergedCell):
                continue
            if c.fill and c.fill.patternType and c.fill.patternType != "none":
                c.fill = NO_FILL

def apply_iizume(ws):
    # 月ごとの全列範囲: 5月=A-F(1-6), 6月=G-L(7-12), 7月=M-R(13-18), 8月=S-X(19-24)
    month_col_range = {5: (1, 6), 6: (7, 12), 7: (13, 18), 8: (19, 24)}
    for month, days in iizume_days_by_month.items():
        c1, c2 = month_col_range[month]
        for day in days:
            row = day + 3
            for col in range(c1, c2 + 1):
                cell = ws.cell(row, col)
                if isinstance(cell, MergedCell):
                    continue
                cell.fill = IIZUME_FILL

for ws_x in (ws1, ws2):
    clear_all_fills(ws_x)
    apply_iizume(ws_x)

# ════════════════════════════════════════════
# リーグ勝敗表 — プロフェッショナルデザイン
# ════════════════════════════════════════════

def sf(h):
    return PatternFill(patternType="solid", fgColor=h)

def mf(size=10, bold=False, color="212121", name="メイリオ"):
    return Font(name=name, size=size, bold=bold, color=color)

AC = Alignment(horizontal="center", vertical="center")

# ─── カラーパレット ─────────────────────
TITLE_BG  = "102A43"  # 深紺: タイトル
HDR_BG    = "1565C0"  # スチールブルー: ヘッダー行4
HDR_BG2   = "0D47A1"  # 濃ブルー: ヘッダー行5 (統計ラベル)
TEAM_ODD  = "1A237E"  # 深インディゴ: 奇数チーム名
TEAM_EVEN = "283593"  # やや薄インディゴ: 偶数チーム名
WL_ODD    = "FFFFFF"  # 白: 奇数チーム勝敗行
SC_ODD    = "F5F5F5"  # 薄グレー: 奇数チームスコア行
WL_EVEN   = "EFF3FF"  # 薄青白: 偶数チーム勝敗行
SC_EVEN   = "E8ECF8"  # 薄青: 偶数チームスコア行
STATS_BG  = "E3F2FD"  # 薄水色: 集計ゾーン (AL-AR)
RANK_BG   = "FFF9C4"  # 薄ゴールド: 順位ゾーン (AS)
INACTIVE  = "EEEEEE"  # 非アクティブ行
WHITE     = "FFFFFF"
GOLD_TXT  = "FFD700"
DARK_TXT  = "212121"
GREY_TXT  = "9E9E9E"
RANK_TXT  = "BF360C"  # 深オレンジ: 順位数字

THIN_G = Side(style="thin",   color="BDBDBD")
MED_N  = Side(style="medium", color="102A43")

ws_r = wb["リーグ勝敗表"]

# ─── タイトル行 (row 1) ─────────────────
c = ws_r.cell(1, 1)
c.fill = sf(TITLE_BG)
c.font = Font(name="メイリオ", size=20, bold=True, color=GOLD_TXT)
c.alignment = AC
ws_r.row_dimensions[1].height = 36

# ─── 細いスペーサー (rows 2-3) ───────────
for r in (2, 3):
    ws_r.row_dimensions[r].height = 4
    for col in range(1, 46):
        cell = ws_r.cell(r, col)
        if isinstance(cell, MergedCell):
            continue
        cell.fill = sf(TITLE_BG)

# ─── ヘッダー行 4 (対戦相手番号) ─────────
for col in range(1, 46):
    cell = ws_r.cell(4, col)
    if isinstance(cell, MergedCell):
        continue
    cell.fill = sf(HDR_BG)
    cell.font = mf(size=11, bold=True, color=WHITE)
    cell.alignment = AC

# ─── ヘッダー行 5 (チーム名・集計ラベル) ──
for col in range(1, 46):
    cell = ws_r.cell(5, col)
    if isinstance(cell, MergedCell):
        continue
    if col >= 38:
        cell.fill = sf(HDR_BG2)
        cell.font = mf(size=12, bold=True, color=WHITE)
    elif col == 45:
        cell.fill = sf(HDR_BG2)
        cell.font = mf(size=12, bold=True, color=GOLD_TXT)
    else:
        cell.fill = sf(HDR_BG)
        cell.font = mf(size=11, bold=True, color=WHITE)
    cell.alignment = AC

# ─── データ行 (rows 6-17: 6チーム×2行) ───
for team_idx in range(6):
    r_wl = 6 + team_idx * 2
    r_sc = r_wl + 1
    is_even = (team_idx % 2 == 1)
    team_bg = TEAM_EVEN if is_even else TEAM_ODD
    wl_bg   = WL_EVEN   if is_even else WL_ODD
    sc_bg   = SC_EVEN   if is_even else SC_ODD

    for r, grid_bg in [(r_wl, wl_bg), (r_sc, sc_bg)]:
        for col in range(1, 46):
            cell = ws_r.cell(r, col)
            if isinstance(cell, MergedCell):
                continue
            if col <= 4:
                cell.fill = sf(team_bg)
                cell.font = mf(size=10, bold=True, color=WHITE)
            elif col == 45:
                cell.fill = sf(RANK_BG)
                cell.font = mf(size=16, bold=True, color=RANK_TXT)
            elif col >= 38:
                cell.fill = sf(STATS_BG)
                cell.font = mf(size=10, bold=(col == 41), color=DARK_TXT)
            else:
                cell.fill = sf(grid_bg)
                cell.font = mf(size=10, color=DARK_TXT)
            cell.alignment = AC

# ─── 非アクティブ行 (rows 18-27) ──────────
for r in range(18, 28):
    for col in range(1, 46):
        cell = ws_r.cell(r, col)
        if isinstance(cell, MergedCell):
            continue
        cell.fill = sf(INACTIVE)
        cell.font = mf(size=9, color=GREY_TXT)
        cell.alignment = AC

# ─── 罫線: 全セル一括 (rows 4-27) ─────────
# ゾーン境界: col 1 (左端), col 5 (対戦グリッド開始=E),
#             col 38 (集計開始=AL), col 45 (順位=AS), col 45 (右端)
ZONE_LEFT_COLS = {1, 5, 38, 45}
PAIR_BOTTOMS   = {6 + i * 2 + 1 for i in range(6)}  # {7,9,11,13,15,17}

for r in range(4, 28):
    for col in range(1, 46):
        cell = ws_r.cell(r, col)
        if isinstance(cell, MergedCell):
            continue
        left   = MED_N if col in ZONE_LEFT_COLS  else THIN_G
        right  = MED_N if col == 45              else THIN_G
        top    = MED_N if r == 4                 else THIN_G
        bottom = MED_N if (r == 5 or r == 27 or r in PAIR_BOTTOMS) else THIN_G
        cell.border = Border(left=left, right=right, top=top, bottom=bottom)

wb.save(DST)
print(f"Saved: {DST}")
